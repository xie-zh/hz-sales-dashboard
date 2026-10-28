# -*- coding: utf-8 -*-
"""
分公司 · 员工三力明细 解析器
================================
读取「分公司-员工三力明细.xlsx」的「三力数据源」表，提取：
  - 每位员工的 三力（门槛破零值 / 标准人力值 / 英雄人力值）
    标准、是否达成、差距
  - 按事业部 / 部门聚合达成人数、考核人数、达成率
供云端快照（_deploy/sanli_data.json）与本地服务使用。

三力口径：
  门槛破零值  -> 是否达成(L)  差距(M)  标准(K)
  标准人力值  -> 是否达成(O)  差距(P)  标准(N)
  英雄人力值  -> 是否达成(R)  差距(S)  标准(Q)
达成判定：标准非空 且 是否达成 == '是'；标准为空 -> 未设标准（不参与达成率）
"""
import os
import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None

import config_loader

# 三力看板排除的状态（离职人员不计入）。改自 config/dashboard.config.json，
# 配置文件缺失时自动回退内置默认值（["离职"]），行为不变。
EXCLUDE_STATUS = set(config_loader.load_config()["sanli"]["exclude_status"])

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SANLI_XLSX = os.path.join(BASE_DIR, "分公司-员工三力明细.xlsx")

FORCES = ["门槛破零值", "标准人力值", "英雄人力值"]

# 三力等级（由高到低），每个员工归入其达到的最高等级
# 英雄 ⊂ 标准 ⊂ 破零；未破零 -> 无效人力
LEVELS = ["英雄人力", "标准人力", "破零人力", "无效人力"]
LEVEL_COLORS = {
    "英雄人力": "#f97316",   # 橙
    "标准人力": "#eab308",   # 黄
    "破零人力": "#3b82f6",   # 蓝
    "无效人力": "#9ca3af",   # 灰
}

# 表头 0 起列号：工号A0 姓名B1 状态C2 分公司D3 事业部E4 部门F5 族类G6 职级H7 岗龄I8
# 利润业绩/销售业绩 J9
FORCE_COLS = {
    "门槛破零值": {"标准": 10, "是否达成": 11, "差距": 12},
    "标准人力值": {"标准": 13, "是否达成": 14, "差距": 15},
    "英雄人力值": {"标准": 16, "是否达成": 17, "差距": 18},
}


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "-", "None", "#REF!", "#REF"):
        return None
    try:
        return float(s.replace(",", "").replace("%", ""))
    except Exception:
        return None


def _s(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _force_state(std, done):
    """返回单个力的状态：达成 / 未达成 / 未设"""
    if std is None:
        return "未设"
    if done == "是":
        return "达成"
    return "未达成"


def _level(sanli):
    """三力等级：英雄 ⊂ 标准 ⊂ 破零，取达成最高等级；未破零 -> 无效人力"""
    if sanli["英雄人力值"]["是否达成"] == "是":
        return "英雄人力"
    if sanli["标准人力值"]["是否达成"] == "是":
        return "标准人力"
    if sanli["门槛破零值"]["是否达成"] == "是":
        return "破零人力"
    return "无效人力"


def parse_employees():
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl，请先安装：pip install openpyxl")
    wb = openpyxl.load_workbook(SANLI_XLSX, read_only=True, data_only=True)
    if "三力数据源" not in wb.sheetnames:
        wb.close()
        raise RuntimeError("未找到「三力数据源」工作表")
    ws = wb["三力数据源"]
    rows = list(ws.iter_rows(values_only=True))
    employees = []
    for r in rows[1:]:
        if not r:
            continue
        name = _s(r[1]) if len(r) > 1 else ""
        if name == "":
            continue
        status = _s(r[2]) if len(r) > 2 else ""
        if status in EXCLUDE_STATUS:  # 排除状态（默认：离职）不计入三力看板
            continue
        emp = {
            "工号": _s(r[0]) if len(r) > 0 else "",
            "姓名": name,
            "状态": status,
            "事业部": _s(r[4]) if len(r) > 4 else "",
            "部门": _s(r[5]) if len(r) > 5 else "",
            "职级": _s(r[7]) if len(r) > 7 else "",
            "岗龄": _s(r[8]) if len(r) > 8 else "",
            "业态": _s(r[19]) if len(r) > 19 else "",
            "业务模式": _s(r[20]) if len(r) > 20 else "",
            "业绩": _num(r[9]) if len(r) > 9 else None,
            "是否有效人力": _s(r[31]) if len(r) > 31 else "",
        }
        sanli = {}
        for f in FORCES:
            c = FORCE_COLS[f]
            std = _num(r[c["标准"]]) if len(r) > c["标准"] else None
            done = _s(r[c["是否达成"]]) if len(r) > c["是否达成"] else ""
            gap = _num(r[c["差距"]]) if len(r) > c["差距"] else None
            sanli[f] = {
                "标准": std,
                "是否达成": done,
                "差距": gap,
                "状态": _force_state(std, done),
            }
        emp["三力"] = sanli
        emp["等级"] = _level(sanli)
        employees.append(emp)
    wb.close()
    return employees


def build_meta():
    employees = parse_employees()
    try:
        mtime = os.path.getmtime(SANLI_XLSX)
        dt = datetime.datetime.fromtimestamp(mtime)
        data_time = "{}月{}日 {:02d}:{:02d}".format(dt.month, dt.day, dt.hour, dt.minute)
        mtime_str = dt.strftime("%Y-%m-%d %H:%M:%S")
    except OSError:
        mtime = 0.0
        data_time = ""
        mtime_str = ""

    # 全局汇总
    summary = {}
    for f in FORCES:
        done = sum(1 for e in employees if e["三力"][f]["状态"] == "达成")
        notdone = sum(1 for e in employees if e["三力"][f]["状态"] == "未达成")
        unset = sum(1 for e in employees if e["三力"][f]["状态"] == "未设")
        assessed = done + notdone
        summary[f] = {
            "达成": done,
            "未达成": notdone,
            "未设": unset,
            "考核": assessed,
            "达成率": (done / assessed) if assessed > 0 else None,
        }
    # 三力等级汇总
    summary["等级"] = {lv: sum(1 for e in employees if e["等级"] == lv)
                       for lv in LEVELS}
    summary["有效人力"] = sum(1 for e in employees if e["是否有效人力"] == "是")

    # 事业部聚合
    bus_map = {}
    dept_map = {}  # (事业部, 部门) -> dict
    for e in employees:
        bu = e["事业部"] or "未知事业部"
        dept = e["部门"] or "未知部门"
        if bu not in bus_map:
            bus_map[bu] = {"事业部": bu, "总人数": 0,
                           "达成": {f: 0 for f in FORCES},
                           "考核": {f: 0 for f in FORCES},
                           "等级": {lv: 0 for lv in LEVELS},
                           "有效人力": 0}
        bus_map[bu]["总人数"] += 1
        bus_map[bu]["等级"][e["等级"]] += 1
        if e["是否有效人力"] == "是":
            bus_map[bu]["有效人力"] += 1
        key = (bu, dept)
        if key not in dept_map:
            dept_map[key] = {"部门": dept, "总人数": 0,
                             "达成": {f: 0 for f in FORCES},
                             "考核": {f: 0 for f in FORCES},
                             "等级": {lv: 0 for lv in LEVELS},
                             "有效人力": 0}
        dept_map[key]["总人数"] += 1
        dept_map[key]["等级"][e["等级"]] += 1
        if e["是否有效人力"] == "是":
            dept_map[key]["有效人力"] += 1
        for f in FORCES:
            st = e["三力"][f]["状态"]
            if st == "未设":
                continue
            bus_map[bu]["考核"][f] += 1
            dept_map[key]["考核"][f] += 1
            if st == "达成":
                bus_map[bu]["达成"][f] += 1
                dept_map[key]["达成"][f] += 1

    bus_list = []
    for bu in sorted(bus_map.keys()):
        b = bus_map[bu]
        b["部门"] = [dept_map[(bu, d)] for d in sorted(
            [k[1] for k in dept_map if k[0] == bu])]
        b["达成率"] = {f: (b["达成"][f] / b["考核"][f] if b["考核"][f] > 0 else None)
                       for f in FORCES}
        bus_list.append(b)

    return {
        "meta": {
            "file_mtime": mtime_str,
            "file_mtime_ts": mtime,
            "data_time": data_time,
            "server_time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "sheet": "三力数据源",
        },
        "total": len(employees),
        "forces": FORCES,
        "levels": LEVELS,
        "level_colors": LEVEL_COLORS,
        "summary": summary,
        "bus": bus_list,
        "employees": employees,
    }


if __name__ == "__main__":
    import json
    data = build_meta()
    print(json.dumps({"meta": data["meta"], "summary": data["summary"],
                      "bus": data["bus"][:2]}, ensure_ascii=False, indent=2))
